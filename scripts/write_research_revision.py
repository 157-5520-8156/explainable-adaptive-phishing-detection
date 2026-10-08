"""Rewrite substantive chapters using genuine retained Cite controls and template components."""
from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile,ZIP_DEFLATED
import json,re,hashlib
import pandas as pd
from lxml import etree as E
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];BUILD=ROOT/'build/redesign_20261007';FINAL=ROOT/'thesis/Phishing_Detection_Mendeley_Draft.docx'
BASE=ROOT/'thesis/archive/Phishing_Detection_Mendeley_20261006.docx'
REF=ROOT/'Project_Proposal_New_Template_Research_Oriented_Projects.docx'
assert hashlib.sha256(REF.read_bytes()).hexdigest()=='e449eedfdd999ef876f98e97d3a023270740b6691a0a1e8203a7ca6dc2d39af2'
doc=Document(BASE);body=doc._element.body;original=list(body)
source=Document(REF);eq_template=deepcopy(source.tables[2]._tbl);data_template=deepcopy(source.tables[3]._tbl)
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','wp':'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing'}
def xp(node,expr):return E._Element.xpath(node,expr,namespaces=NS)
def text(node):return ''.join(xp(node,'.//w:t/text()'))
def tag(x):
    node=x.find(qn("w:sdtPr")); node=node.find(qn("w:tag")) if node is not None else None
    return node.get(qn("w:val"),"") if node is not None else ""
def controls(node):return [deepcopy(x) for x in xp(node,".//w:sdt") if "MENDELEY_CITATION" in tag(x)]
cites={i:controls(x) for i,x in enumerate(original)}
old_ids=[tag(x) for x in xp(body,'.//w:sdt') if 'MENDELEY_CITATION' in tag(x)]
blocks=[]
def font(r,size=11,bold=False):r.font.name='Arial';r.font.size=Pt(size);r.font.bold=bold;r.font.color.rgb=RGBColor(0,0,0)
def fill(p,content,citation_nodes=None):
    for x in list(p._p):
        if x.tag!=qn('w:pPr'):p._p.remove(x)
    for token in re.split(r'(\{C\d+\}|\{I\d+\})',content):
        if re.fullmatch(r'\{C\d+\}',token):
            idx=int(token[2:-1]);p._p.append(cites[idx].pop(0))
        elif re.fullmatch(r'\{I\d+\}',token):
            idx=int(token[2:-1]);p._p.append(citation_nodes[idx])
        else:font(p.add_run(token))
    p.paragraph_format.line_spacing=1.5;p.paragraph_format.space_after=Pt(0)
    p.paragraph_format.first_line_indent=Pt(24)
    p.paragraph_format.keep_with_next=False;p.paragraph_format.page_break_before=False
    return p

def replace(index,content):
    nodes=cites[index];p=Paragraph(original[index],doc._body);fill(p,content,nodes)
# Cover revision date only; metadata and source logo stay intact.
for node in original[:14]:
    for t in xp(node,'.//w:t'):
        if t.text and '6 October 2026' in t.text:t.text=t.text.replace('6 October 2026','7 October 2026')
replace(15,"This study develops and evaluates an offline explainable phishing detector and label-aware update policies using three public datasets. An earlier URL forest achieved source F1 0.9963 but nearly all legitimate external URLs were falsely flagged. An audit identifies collection shortcuts and a root-path representation defect. A redesigned URL representation suppresses scheme and leading www cues, preserves meaningful path case and enforces domain isolation; however, its frozen cross-source result remains inadequate. A separate archived-page branch combines 44 recomputed URL statistics with 24 published content measurements, excluding external reputation features. Histogram gradient boosting is selected using calibration domains only. On 6,849 untouched-domain records it achieves phishing recall 0.8222, precision 0.9293, F1 0.8725, average precision 0.9528 and legitimate false-positive rate 0.0624. Tree-path-dependent SHAP reconstructs the selected model's log-odds output with maximum error 8.44e-15 across 200 records. Forty-eight controlled replays compare fixed, threshold-only, periodic and ADWIN policies under two label delays and two orders. Periodic updating with two-batch delay in the path-profile order reduces mean final false-positive rate to 0.0495 while recall falls to 0.7959. Adaptation is therefore a measured tradeoff, not an unconditional improvement. The implementation retains predictions, label-arrival ledgers, checkpoints, source snapshots and independent calculation checks. The results support an inspectable archived-page prototype while leaving cross-source content generalisation, live extraction and natural-time evaluation unresolved.")
replace(16,"Keywords: phishing detection; archived webpage features; URL normalization; SHAP; domain isolation; adaptive retraining; delayed labels.")
replace(67,"The study begins with PhiUSIIL, which supplies URL strings, published labels and precomputed website attributes {I0}. Its original framework includes incremental learning in its scope {I1}. A second URL corpus published by Wangchuk supports a frozen transfer test {I2}. The project then adds the Hannousse and Yahiouche archive to study a separate model that uses both URL strings and measured webpage content {NEW-DATA}. These inputs require different contracts: a URL-only detector needs a string, whereas the archived-page detector also needs 24 content measurements. The experiments and claims distinguish those contracts.")
# Original paragraph had four? Assert and use precise nodes count below, not silently discard.
replace(70,"An apparently strong source result can arise when train and test samples share collection conventions even after registered domains have been separated. All retained legitimate PhiUSIIL records use HTTPS and empty paths; its validation legitimate records also all have a www prefix. Another corpus does not share those conventions. The first model therefore requires diagnosis and redesign before it can support a usable prototype.")
replace(71,"The research questions are: RQ1, what do domain-isolated and cross-source tests reveal about URL-only phishing classification, and what detection performance is obtained when archived webpage measurements are available? RQ2, do verified SHAP explanations and feature-group diagnostics expose the basis and errors of the selected detector? RQ3, how do fixed, threshold-only, periodic and drift-triggered policies compare when labels arrive immediately or after two batches, and how much adaptation occurs in a shuffled control?")
replace(73,"The project aim is to implement an explainable phishing classification prototype with explicit input requirements and adaptive update policies, and to establish its measured performance and limits using reproducible archived data.")
replace(75,"The objectives are to audit public data and isolate evaluation domains; compare regularized linear and tree models using development data; identify and repair representation shortcuts; evaluate a separate archived-page detector when URL strings alone are insufficient; verify the selected model's explanations; and compare update policies using identical replay records and label schedules.")
replace(76,"A further objective is to provide a runnable prediction and explanation interface together with dataset hashes, split manifests, fitted checkpoints, per-record predictions and independent verification. A successful pipeline must report false-positive rate and phishing recall together. Threshold calibration alone cannot establish detector quality.")
replace(78,"The pipeline has two revised branches. The URL branch combines development records from PhiUSIIL and Wangchuk and reserves a third corpus for transfer assessment. The archived-page branch uses domain-isolated portions of the third corpus for initial fitting, calibration, controlled feedback and a final evaluation. Its results concern unseen domains from that corpus, rather than unseen-source content generalisation. The failed original experiment remains an auditable baseline.")
replace(80,"The task is binary classification of archived records with phishing positive. URL features are recomputed locally. Page features are the creator's released measurements; the project does not independently re-extract them from HTML. External reputation, blacklist, WHOIS, DNS, traffic, page-rank and Google-index attributes are excluded from the revised page model. No listed destination is visited, and published labels are not claims about current website behaviour.")
replace(81,"The replay uses actual archived records and unchanged labels. Record order and delay are controlled experimental choices because per-record observation and verification timestamps are unavailable. Domain ordering by average path length produces a constructed covariate profile; a shuffled version of exactly the same records is the stationary-order control. Neither experiment establishes naturally occurring chronological drift.")
replace(83,"The report is intended for academic assessment and for developers examining an offline phishing prototype. The program can classify and explain records that satisfy its input contract. Its measurements do not justify deployment on arbitrary current websites; the remaining false alarms, missed phishing and unavailable live extraction are explicit limits.")
# Update relevant literature text while preserving all citation controls already in place.
replace(100,"The revised page classifier has correlated URL and content inputs. Its accepted explanations use TreeSHAP's training-path formulation and describe contributions to log odds. This differs from interventional explanations with an explicit background distribution. Both formulations depend on the fitted model and the chosen treatment of feature dependence. Numerical reconstruction is required for an explanation to enter the results; it does not establish a causal security rule or improved human trust.")
replace(108,"The research position is an integration and evaluation study: collection-aware URL representation, an archived-page prototype, verified model explanations and label-aware adaptation are examined with separate evidence boundaries. Existing classifier, attribution and drift algorithms are used rather than claimed as new inventions. Adding content changes the available information and the training corpus, so its benefit requires within-partition diagnostics rather than an unqualified comparison with the failed URL baseline.")
# Add primary-source discussion before Chapter 3, without copying unverified prior performance.
lit=doc.add_paragraph(style='Normal');fill(lit,"Hannousse and Yahiouche describe a benchmark construction method and compare URL, content and external-service feature groups {NEW-PAPER}. Their archive supplies a way to study historical phishing pages without fetching present destinations. The present study excludes external-service features, uses its own URL extractor and groups by registered domain. These choices make the new results a different protocol; the authors' headline accuracy is not used as a directly comparable result.")
body.remove(lit._p);body.insert(body.index(original[114]),lit._p)
# Cache section break components and remove only Chapters 3-5. References remain real Cite content.
section_breaks={3:deepcopy(original[184]),4:deepcopy(original[241]),5:deepcopy(original[253])}
insert_at=body.index(original[115])
for x in original[115:254]:body.remove(x)
reference_heading=original[254]
new_nodes=[]
def attach(node):
    if node.getparent() is body:body.remove(node)
    body.insert(body.index(reference_heading),node);new_nodes.append(node)
def p(s):
    x=doc.add_paragraph(style='Normal');fill(x,s);attach(x._p);blocks.append({'type':'paragraph','text':s});return x
counter=1000
heads=[]
def h(s,level=2):
    if re.match(r'^\d+\.\d+\.\d+ ',s):level=3
    global counter
    x=doc.add_paragraph(s,style='Heading '+str(level));x.paragraph_format.page_break_before=False
    x.paragraph_format.first_line_indent=Pt(0);x.paragraph_format.line_spacing=1.5;x.paragraph_format.keep_with_next=True
    x.alignment=WD_ALIGN_PARAGRAPH.CENTER if level in [1,3] else WD_ALIGN_PARAGRAPH.LEFT
    for r in x.runs:font(r,18 if level==1 else 15,True)
    b=OxmlElement('w:bookmarkStart');b.set(qn('w:id'),str(counter));b.set(qn('w:name'),f'research_h_{counter}');x._p.insert(1,b)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(counter));x._p.append(end);counter+=1
    attach(x._p);heads.append(s);blocks.append({'type':'heading','text':s,'level':level})
def strip(node):
    for x in list(xp(node,'.//w:bookmarkStart|.//w:bookmarkEnd|.//w:sdt')):x.getparent().remove(x)
def table(caption,headers,rows,widths):
    cap=p(caption);cap.style='题注1';cap.paragraph_format.first_line_indent=Pt(0);cap.alignment=WD_ALIGN_PARAGRAPH.CENTER;cap.paragraph_format.keep_with_next=True
    template=deepcopy(data_template);strip(template);prototypes=[deepcopy(x) for x in template.findall(qn('w:tr'))]
    for x in list(template.findall(qn('w:tr'))):template.remove(x)
    grid=template.find(qn('w:tblGrid'))
    for x in list(grid):grid.remove(x)
    for width in widths:
        col=OxmlElement('w:gridCol');col.set(qn('w:w'),str(round(width*1440)));grid.append(col)
    for i,values in enumerate([headers]+rows):
        tr=deepcopy(prototypes[0 if i==0 else min(1,len(prototypes)-1)]);cells=[deepcopy(x) for x in tr.findall(qn('w:tc'))]
        for x in list(tr.findall(qn('w:tc'))):tr.remove(x)
        for j,(value,width) in enumerate(zip(values,widths)):
            tc=deepcopy(cells[min(j,len(cells)-1)]);pr=tc.find(qn('w:tcPr'))
            for x in list(tc):
                if x is not pr:tc.remove(x)
            pr.find(qn('w:tcW')).set(qn('w:w'),str(round(width*1440)))
            for x in list(pr):
                if x.tag in [qn('w:tcBorders'),qn('w:shd')]:pr.remove(x)
            borders=OxmlElement('w:tcBorders')
            for side in ['top','bottom','left','right']:
                x=OxmlElement('w:'+side);x.set(qn('w:val'),'single' if side=='top' or (side=='bottom' and i==len(rows)) else 'none');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'000000');borders.append(x)
            pr.append(borders);tc.append(OxmlElement('w:p'));tr.append(tc)
        template.append(tr)
    attach(template);t=Table(template,doc._body);t.autofit=False
    for i,(row,values) in enumerate(zip(t.rows,[headers]+rows)):
        pr=row._tr.get_or_add_trPr()
        for x in list(pr.findall(qn('w:trHeight'))):pr.remove(x)
        pr.append(OxmlElement('w:cantSplit'))
        if i==0:pr.append(OxmlElement('w:tblHeader'))
        for cell,value in zip(row.cells,values):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;x=cell.paragraphs[0]
            x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.first_line_indent=Pt(0);x.paragraph_format.line_spacing=1.5;x.paragraph_format.space_after=Pt(0)
            x.paragraph_format.keep_with_next=len(rows)<=6 and i<len(rows)
            font(x.add_run(str(value)),10,i==0)
    blocks.append({'type':'table','caption':caption,'headers':headers,'rows':rows})
    p('').paragraph_format.first_line_indent=Pt(0)
def eq(number):
    template=deepcopy(eq_template);strip(template);attach(template);t=Table(template,doc._body);t.autofit=False
    for cell,width in zip(t.rows[0].cells,[5.13,.57]):
        cell.width=Inches(width);cell.text='';x=cell.paragraphs[0];x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.first_line_indent=Pt(0);x.paragraph_format.line_spacing=1.5;x.paragraph_format.space_after=Pt(4)
    math=E.parse(str(ROOT/'thesis/equations'/f'eq_{number}.omml')).getroot()
    if math.tag!=qn('m:oMath'):math=math.find('.//'+qn('m:oMath'))
    t.rows[0].cells[0].paragraphs[0]._p.append(deepcopy(math));font(t.rows[0].cells[1].paragraphs[0].add_run('('+number+')'))
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'));blocks.append({'type':'equation','number':number});p('')
def fig(name,caption,width=5.65):
    x=doc.add_paragraph();x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.first_line_indent=Pt(0);x.paragraph_format.line_spacing=1;x.paragraph_format.keep_with_next=True
    x.add_run().add_picture(str(ROOT/'experiments/page_20261007_02/analysis/figures'/name),width=Inches(width));attach(x._p)
    c=p(caption);c.style='题注1';c.alignment=WD_ALIGN_PARAGRAPH.CENTER;c.paragraph_format.first_line_indent=Pt(0);blocks.append({'type':'figure','caption':caption,'path':name})
def endchapter(n):attach(section_breaks[n])
def pct(v):return f'{100*float(v):.2f}'
def val(v):return f'{float(v):.4f}'

h('Chapter 3. Methodology',1)
h('3.1 Problem Statement')
p("The response variable y is 1 for published phishing records and 0 for legitimate records. Two input contracts are investigated. The revised URL classifier takes a normalized learning view of a URL, using either 44 numerical statistics or character n-gram TF-IDF. The archived-page classifier takes those numerical statistics and 24 released webpage-content measurements. The two branches answer different questions and their results are not pooled into a single accuracy figure.")
p("For each contract the fitted estimator produces a phishing score, and a separately selected cutoff produces a binary decision. Scores are not treated as calibrated real-world probabilities: balanced or domain-weighted fitting changes the represented population. The main measurements are false-positive rate, recall, precision, F1, average precision and domain-macro rates. Both a low false-positive rate and useful phishing recall are needed; always predicting legitimate is retained as a counterexample.")
h('3.2 Approach')
h('3.2.1 Audit and Redesign')
p("The original archived run remains unchanged. A later audit checks raw-file hashes, label mappings, class order, feature names and feature recomputation. It also compares root URLs with and without a terminal slash. The redesign corrects this representation and suppresses the observed source-specific scheme and www patterns. These development decisions are made before the new corpus test is scored. Earlier candidate runs, including a character-score shape error, remain identifiable rather than being silently replaced.")
p("The URL-only development branch still has weak recall at strict false-alarm cutoffs. This motivates a separately documented page-feature branch before its test scoring. It is a change in scope and available information, not a claim that normalizing a URL alone repairs the whole task. The selected page estimator and all replay policies are frozen before final evaluation. Feature-group ablations added after primary test inspection are labelled exploratory and cannot redefine the primary selection.")
h('3.2.2 Classification Models')
p("The URL branch compares numerical Logistic Regression, Random Forest, histogram gradient boosting and character Logistic Regression. Logistic Regression uses a standardised numerical matrix or a sparse TF-IDF matrix with C=1. The TF-IDF vocabulary is fitted only on training rows, uses character lengths 3-5, min_df=3 and at most 80,000 features. A PhiUSIIL-only forest provides a source-only comparator and does not enter mixed-source model selection. Equation (3-1) defines the logistic score.")
eq('3-1')
p("The revised mixed-source forest uses 200 trees, depth limit 20, minimum leaf size 4 and square-root feature sampling. Histogram boosting uses 180 iterations, 31 leaves, learning rate 0.08 and L2 regularisation 1.0, without early stopping. Fitting weights first equalise source/class mass and then registered-domain contribution within each stratum. A training cap of 30,000 rows per source/class and 12 rows per source/class/domain limits collection concentration. Equation (3-2) gives the forest's probability averaging rule.")
eq('3-2')
p("The archived-page branch compares a 200-tree forest with minimum leaf size 2 against histogram gradient boosting with 180 iterations, 15 leaves, learning rate 0.08 and L2 regularisation 2.0. Both use the same 68 inputs, fitting rows, calibration rows and class/domain weighting. The candidate configurations are specified before testing. Fitting seeds are 17, 42 and 73. Histogram boosting is deterministic here: identical initial scores across those seeds are reported as identical executions, not as three independent samples of uncertainty.")
h('3.2.3 Evaluation and Cutoffs')
p("Equation (3-3) defines precision P, phishing recall R, F1 and legitimate false-positive rate FPR from the four confusion-matrix counts. These rates have different denominators. Domain-macro FPR first computes the false-positive fraction within every domain containing legitimate records and then averages domains equally; domain-macro recall is defined analogously over phishing-containing domains. A mixed-label domain may contribute to both measures.")
eq('3-3')
p("Average precision measures ranking across the available score cutoffs rather than one operating point. Equation (3-4) uses successive recall increments and their corresponding precision. This is the noninterpolated definition used by scikit-learn {C129}. Tied scores are grouped. Figure 4-1 shows the empirical precision-recall step curve, the test prevalence baseline and the operating point fixed on calibration records.")
eq('3-4')
p("For a requested observed FPR budget, normal calibration scores are sorted from highest to lowest. If k=floor(alpha times normal count), the cutoff is the next representable floating-point value above the normal score at zero-based position k. This excludes ties as a group. The mixed-source cutoff is the maximum required across represented sources. The primary URL budget is 1%, with 5% reported as a secondary point. The separate page branch uses 5% as primary and also records 1%. These are empirical calibration budgets, not population guarantees.")
p("Model-family selection uses mean calibration phishing recall at the primary budget, with average precision as the tie breaker. The URL ranking averages sources and fitting seeds; the page ranking averages fitting seeds. No internal test, transfer test or final checkpoint test chooses the model or cutoff. After those tests have been inspected, further development using them must be reported as exploratory or assessed on another genuinely untouched sample.")
h('3.2.4 Verified Explanations')
p("Equations (3-5) and (3-6) state the additive attribution and Shapley formulation {C132}. For the revised page model M=68 and f is the histogram-boosting raw log-odds score. The logistic transform of the reconstructed log odds must match the phishing score. Positive contributions increase that score relative to the explanation baseline; they do not independently establish that a website is phishing.")
eq('3-5');eq('3-6')
p("Accepted explanations use TreeSHAP with feature_perturbation=tree_path_dependent and model_output=raw. The background information consists of training-path counts in the fitted trees. Two hundred test records are sampled with seed 20261007 for global inspection; the first saved example of each confusion-matrix category supplies four local diagnostics. Numerical additivity is checked against the estimator's actual decision_function. A different interventional configuration with an explicit 100-row background failed this check and is retained as rejected evidence rather than interpreted.")
p("Post hoc URL-only and content-only ablations use the same page-branch fitting, calibration and test partitions as the primary combined model. Their cutoffs are independently calibrated at the same 5% budget. They diagnose the complementarity of feature groups within one fixed partition; their seed-42 measurements are not confirmatory evidence for a newly selected architecture. No human-subject study, explanation usability score or causal intervention is claimed.")
h('3.2.5 Adaptive Policies and Label Timing')
p("Fixed prediction retains its initial estimator and cutoff. Threshold-only updating recalibrates the original estimator whenever labels arrive. Periodic updating refits and recalibrates every four arrival batches. Drift-triggered updating monitors original prediction errors with ADWIN delta=0.002 and allows a refit after an alarm and a two-batch cooldown. ADWIN supplies a change signal, not an independent diagnosis of its cause {C137}. All policies receive the same records within an order/seed/delay realization.")
eq('3-7');eq('3-8')
p("In Equation (3-7), b(i) is the prediction batch and d is zero or two. A batch is predicted before any of its newly available labels may update the decision. Arrival records preserve the originating decision version. A version includes both the estimator and cutoff. Any update increments that version and resets ADWIN; labels for predictions under an older version remain available for fitting, but their errors are excluded from the reset detector's state. Their original predictions still count in replay metrics.")
p("The update pool combines at most 800 fixed initial anchors with the latest 1,600 revealed replay rows. Normalized URL duplicates are removed. GroupShuffleSplit separates update fitting and calibration domains with calibration fraction 0.30 and a seed derived from the realization and batch {C158}. Both parts must contain both classes. The threshold-only policy leaves the estimator fixed. Labels still pending when the stream ends are not flushed solely to improve a final checkpoint.")
p("The replay uses the 2,283 rows from the reserved feedback fold in batches of 160, giving fourteen full batches and one partial batch. In path-profile order, domains are sorted by their mean URL path length, with seed-dependent tie ordering; records within a domain remain adjacent. The stationary-order control shuffles exactly the same pool. Four policies, three seeds, two delays and two orders give 48 executions. Path-profile order is a constructed covariate challenge, and the finite shuffled control is not proof that every underlying distribution is stationary.")
h('3.3 Technologies and Data')
h('3.3.1 Dataset Provenance and Quality')
p("The original PhiUSIIL download has 235,795 rows and the earlier eligibility filter retains 235,152. Raw label 1, legitimate, is mapped to internal 0; raw label 0, phishing, is mapped to internal 1 {C147}. The redesign operates on the retained raw URL/label pairs and removes additional normalized duplicates and conflicts. The Wangchuk raw file has 149,726 rows, with published label 0 legitimate and label 1 phishing {C153}. Only its previously designated 37,392-row adaptation partition enters revised development; its already inspected holdout remains reused diagnostic evidence.")
p("The Hannousse V3 archive contains 11,430 balanced rows collected in May 2020 and released with URL strings, labels and feature measurements. Its creator record and associated paper are primary provenance sources. The common URL quality policy excludes internal whitespace, angle brackets and double quotes, requires an HTTP(S) hostname and preserves original published labels. Four phishing rows fail eligibility and eleven duplicate normalized URLs are removed, leaving 11,415 rows. Raw files, cleaned row identities and excluded counts remain in the execution archive.")
table('Table 3-1. Public raw data and their role in the revised study.', ['Source','Raw rows','Available inputs','Role'], [['PhiUSIIL','235,795','URLs and supplied attributes','URL development'],['Wangchuk V1','149,726','URLs and labels','Reserved URL adaptation partition'],['Hannousse V3','11,430','URLs and archived feature measurements','Transfer test and separate page study']], [1.15,.75,1.8,2.0])
p("SHA-256 checksums identify the exact downloaded bytes: PhiUSIIL a236549cd369cd80bd478ff8e1779cbf44c58d5c3f79f7a51a1adbed7d06d1c6, Wangchuk fccce1a99282b3cd8a1fe08fa83350fe44f7afc4c8889a6d0a94ef9c4d7897e3, and Hannousse 21093e2902e5441c86a6daf95e86e7c332046e477fdf109a579d7bd81e586d6c. These fingerprints identify data; they do not verify the truth of historical labels. The two later archives are used under their published CC BY 4.0 terms.")
h('3.3.2 Representation and Feature Contract')
p("Canonical identity lowercases the scheme and hostname, preserves user information and meaningful path/query/fragment case, and represents an empty root path by a slash. RFC 9110 Section 4.2.3 describes empty-path and slash normalization for HTTP(S) URIs, apart from the OPTIONS exception {NEW-RFC}. This rule applies to the string representation, without making requests. It does not collapse different pages, different hosts, default ports or percent-encoded paths into one identity.")
p("The learning projection is deliberately different from canonical identity: it fixes the scheme to http and suppresses a leading www when the hostname has at least two dots. This removes two observed collection shortcuts without asserting that different schemes or www/non-www destinations are the same resource. Deduplication, domain grouping and published labels use canonical identities rather than the reduced feature view. Feature names representation_length and represented_host_length signal that lengths refer to that view.")
p("The 44 numerical inputs cover string lengths, delimiter counts, character ratios, entropy, hostname and path words, longest runs, path extension length and explicit-port presence. The has_https column is excluded. No target or supplied reputation feature enters that matrix. Training-only standardisation and TF-IDF fitting prevent information from validation or test records entering their learned transformations. The implementation follows scikit-learn's separation of fitting and evaluation {C149}.")
p("The page branch adds exactly the 24 released content columns listed in Table 3-2. Values are joined by unique canonical URL with a one-to-one check and must be finite. They retain the creator's units; a name beginning ratio does not guarantee a common 0-1 scale. External-service attributes and the released URL blacklist/statistical-report fields are excluded. The program rejects an input record missing any content measurement rather than inventing or filling webpage evidence.")
table('Table 3-2. The 24 supplied webpage-content inputs.', ['Group','Released feature names'], [['Links and resources','nb_hyperlinks; ratio_intHyperlinks; ratio_extHyperlinks; ratio_nullHyperlinks; nb_extCSS; links_in_tags'],['Redirects and errors','ratio_intRedirection; ratio_extRedirection; ratio_intErrors; ratio_extErrors'],['Forms and media','login_form; external_favicon; submit_email; ratio_intMedia; ratio_extMedia; sfh'],['Interface and text','iframe; popup_window; safe_anchor; onmouseover; right_clic; empty_title; domain_in_title; domain_with_copyright']], [1.25,4.45])
h('3.3.3 Domain Isolation')
p("Registered domains are obtained using the local domain-extraction configuration and public-suffix rules recorded in the archive. A single five-fold StratifiedGroupKFold assignment fixes the Hannousse roles with seed 20261007 {C160}. Folds 0-2 reserve 6,849 rows for evaluation; fold 3 supplies 2,283 initial rows; fold 4 supplies 2,283 feedback rows. The initial fold is then split by domain into 1,728 fitting rows and 555 calibration rows. These are role partitions, not a fivefold cross-validation estimate.")
p("For the revised URL branch, all domains appearing anywhere in retained Hannousse are removed from both development sources before their grouped split. This removes 4,775 PhiUSIIL and 4,894 Wangchuk rows. Global development grouping keeps a domain in only one role even when it appears in both sources. Training uses folds 0-2, validation fold 3 and internal test fold 4. The actual capped fitting sample has 68,246 rows; row and domain counts are retained per source and class.")
counts=json.loads((ROOT/'experiments/page_20261007_02/artifacts/partition_counts.json').read_text())
table('Table 3-3. Mutually domain-disjoint roles for the archived-page branch.', ['Role','Rows','Legitimate','Phishing','Domains'], [[k.replace('_',' ').title(),v['n'],v['normal'],v['phishing'],v['domains']] for k,v in counts.items()], [1.9,.9,.95,.9,1.05])
p("The same Hannousse test domains serve two explicitly distinguished assessments. For the URL classifier trained only on the other corpora, they represent an unseen source with all domains excluded from development. For the page classifier trained on Hannousse fold 3, they represent new domains within the same corpus. After either test has been read, the records cannot be described as a fresh blind test for subsequent redesign. Common public phishing feeds may still connect campaigns across nominally separate corpora.")
h('3.3.4 Uncertainty and Implementation Environment')
p("A whole-domain bootstrap with 1,000 replicates and fixed seed 20261007 estimates 95% percentile intervals conditional on a saved fitted model. Each replicate samples registered-domain clusters with replacement and retains all records in each sampled cluster. The paired periodic-minus-fixed comparison uses the same sampled clusters for both predictions. Its intervals concern this corpus and these trained decisions; they exclude label error, corpus-selection uncertainty and alternative split realizations.")
p("Python executes the offline experiments with pinned NumPy, pandas, scikit-learn, SciPy, SHAP and River versions. The final record interface validates the 68-feature contract, loads a locally produced checkpoint, reports its cutoff and decision version, and exports all signed SHAP contributions. It does not retrieve website content. Model files use joblib and are treated as trusted project artifacts rather than a format for accepting arbitrary external files.")
h('3.4 Version Management')
p("The accepted earlier experiment, failure audit, URL redesign candidates and page study have separate directories. Immutable source/configuration snapshots identify their implementation. Each fitted checkpoint, row manifest and saved prediction file is hashed. The page study separates selection_frozen from all_replays_frozen and then test_scored states. Explanatory diagnostics are recorded after primary scoring and retain that status. A rejected interventional SHAP check is stored alongside the accepted explanation audit.")
p("Independent verification reads saved prediction files and recomputes confusion matrices, rather than trusting summary metrics alone. It checks domain overlap, decision comparisons, label schedules, update-time availability and fitting/calibration separation for all 48 replays. Small behaviour tests cover root normalization, meaningful path case, tied-score calibration, missing webpage inputs and delayed errors after a cutoff change. These checks verify named implementation properties; they do not certify deployment or source labels.")
h('3.5 Summary')
p("The protocol connects a failed URL baseline to an explicit redesign and a usable archived-page input contract. It separates fitting, calibration, replay and final evaluation domains, checks explanation reconstruction and records the costs and consequences of updating. The next chapter reports every policy family and explains both effective detection and remaining errors.")
endchapter(3)

h('Chapter 4. Implementation',1)
h('4.1 Implemented Pipeline and Verification')
p("The project now contains raw-data provenance, normalized feature extraction, fitted URL candidates, the archived-page classifier, label-aware policy replay, numerical explanations and a command-line record interface. The archived-page interface takes a JSON record containing a URL and all 24 published content measurements, together with a selected checkpoint. It returns the phishing score, binary decision, cutoff, decision version, baseline log odds and signed contributions for all 68 features. No webpage measurement is inferred from the URL alone.")
p("A saved real record is processed end to end through this public interface, and the transformed explanation reconstructs its model score within tolerance. Independent verification passes for 96 prediction files: 48 original replay ledgers and 48 final-domain checkpoint evaluations. It independently recomputes the four confusion counts and verifies label arrival and domain boundaries. Nine behaviour tests pass. A passing calculation audit establishes consistency of the retained files rather than correctness of every published label.")
h('4.2 Diagnosis of the Earlier URL Failure')
p("The retained original source forest gives mean F1 0.9963 within PhiUSIIL but external F1 0.4359 and legitimate false-positive rate 0.9973. The earlier 40,810-row Wangchuk holdout has been inspected and is now diagnostic evidence. A fresh execution of one saved primary forest flags 29,139 of its 29,230 legitimate rows, giving FPR 0.9969. The slight difference from the earlier mean reflects the particular saved model; it is not a corrected detector.")
p("Raw hashes, label mapping, feature order and recomputed feature values pass the diagnosis. The source convention is the substantive problem. All retained legitimate source URLs use HTTPS and empty paths. In a controlled check of 200 actual source-test legitimate root URLs, the old model correctly labels the originals but flags all 200 after a single root slash is appended. The redesigned URL representation gives exactly equal scores before and after that change on 200 real legitimate roots. Equality is a normalization property, not a detection-rate claim.")
p("Normalization removes a demonstrated defect but does not establish transfer. The revised URL branch suppresses protocol and www cues, balances source/domain contributions and compares four model families. Its development-selected histogram booster remains weak at the strict primary point. At the secondary 5% calibration point, its internal Wangchuk FPR rises to 17.43%, despite validation FPR being below 5%. This gap indicates that a concentrated-domain calibration sample does not guarantee test-domain false-alarm control.")
u=pd.read_csv(ROOT/'experiments/redesign_20261007_03/artifacts/test_metrics.csv')
primary=u[(u.family=='hist_gradient_boosting')&(u.seed==42)]
table('Table 4-1. Frozen revised URL model at its registered calibration budgets.', ['Assessment','Budget %','FPR %','Recall %','F1'], [[r.source+' '+('transfer' if r.evaluation=='unseen_source' else 'internal'),pct(r.budget),pct(r.fpr),pct(r.recall),val(r.f1)] for _,r in primary.iterrows()], [2.0,.85,.95,.95,.95])
p("On Hannousse domains excluded from all URL development, the revised URL booster at the secondary point has FPR 20.50% and recall 60.88%. This remains unsuitable for broad warnings. The URL redesign is therefore retained as a limitation and a diagnostic contribution; it is not the final quality result. The following page branch adds measured information and a separate fitting population, and its results are labelled accordingly.")
h('4.3 Archived Page Model Selection')
v=pd.read_csv(ROOT/'experiments/page_20261007_02/artifacts/validation_metrics.csv');v=v[v.budget.eq(.05)]
table('Table 4-2. Calibration-domain results for the archived-page candidates at the 5% budget.', ['Model','Seed','FPR %','Recall %','AP'], [[r.family.replace('_',' '),int(r.seed),pct(r.fpr),pct(r.recall),val(r.average_precision)] for _,r in v.iterrows()], [2.0,.6,.9,1.0,1.2])
p("Histogram gradient boosting has calibration recall 77.97% at observed FPR 4.70% and is selected before any page-test predictions. The forest recalls range from 72.03% to 76.27% at the same false-positive fraction. The three boosting runs produce the same initial fitted decisions because this configuration is deterministic. Consequently, the reported initial test score is one partition/model result; the three fitting seeds must not be mistaken for independent corroboration.")
h('4.4 Untouched Domain Detection Results')
a=json.loads((ROOT/'experiments/page_20261007_02/analysis/summary.json').read_text());st=a['static_test']
table('Table 4-3. Selected fixed archived-page model on 6,849 untouched-domain records.', ['Measure','Value'], [['False-positive rate',pct(st['fpr'])+'%'],['Phishing recall',pct(st['recall'])+'%'],['Precision',pct(st['precision'])+'%'],['F1',val(st['f1'])],['Average precision',val(st['average_precision'])],['ROC AUC',val(st['roc_auc'])],['Domain-macro FPR',pct(st['domain_macro_fpr'])+'%'],['Domain-macro recall',pct(st['domain_macro_recall'])+'%']], [3.6,2.1])
p("The selected cutoff is 0.714698, fixed using calibration rows. On 3,429 legitimate and 3,420 phishing test records, it produces 214 false positives and 608 false negatives. Recall is 82.22%, precision 92.93%, F1 0.8725 and average precision 0.9528. The FPR of 6.24% is substantially below the failed baseline but exceeds the page study's 5% target. The baseline used a different corpus and input contract, so the magnitude of that reduction is descriptive rather than a paired model-effect estimate.")
fig('figure_4_1_page_pr.png','Figure 4-1. Empirical precision-recall step curve on archived-page test domains. The dot uses the frozen calibration cutoff; the dashed line is test phishing prevalence.')
p("The curve in Figure 4-1 is computed directly from retained scores across their actual cutoffs. It is not a smoothed illustrative success curve. Average precision describes useful ranking, while the operating-point dot shows the actual precision/recall pair at the registered cutoff. Selecting a new dot from the test curve would be post hoc tuning and is not performed.")
fig('figure_4_3_confusion.png','Figure 4-2. Confusion matrix of the fixed archived-page model. Rows are published labels and columns are original model decisions.',4.65)
intervals=json.loads((ROOT/'experiments/page_20261007_02/artifacts/test_cluster_intervals.json').read_text());ci=next(x['cluster_bootstrap_95'] for x in intervals if x['order']=='path_profile' and x['policy']=='fixed' and x['delay']==2)
table('Table 4-4. Domain-cluster uncertainty for the fixed seed-42 model.', ['Measure','Estimate','95% interval'], [['FPR %',pct(st['fpr']),pct(ci['fpr'][0])+' to '+pct(ci['fpr'][1])],['Recall %',pct(st['recall']),pct(ci['recall'][0])+' to '+pct(ci['recall'][1])],['F1',val(st['f1']),val(ci['f1'][0])+' to '+val(ci['f1'][1])]], [1.9,1.4,2.4])
p("Domains rather than individual rows are resampled because records within a domain can share webpage structure and campaign characteristics. The row FPR of 6.24% and domain-macro FPR of 5.08% show that row weighting and domain weighting are not interchangeable. An always-legitimate baseline has zero false positives but zero phishing recall and F1. An always-phishing baseline has 100% FPR and approximately 0.666 F1 in this nearly balanced test, below the selected model.")
h('4.5 Explanations and Feature Group Diagnostics')
ex=json.loads((ROOT/'experiments/page_20261007_02/analysis/explanation_audit.json').read_text())
p("Accepted TreeSHAP explanations reconstruct the boosting raw score with maximum absolute error 8.44e-15 across 200 sampled test rows. Applying the logistic transform also matches predict_proba. The interventional attempt with a 100-row background fails reconstruction for at least one record: 4.920196 versus actual log odds 5.463303. Those attributions are rejected. The accepted training-path configuration has a different treatment of feature dependence; a numerical pass does not make it a causal explanation.")
fig('figure_4_2_shap.png','Figure 4-3. Mean absolute accepted TreeSHAP contributions for the ten largest inputs across 200 test rows. The output unit is log odds.')
p("The largest mean absolute contribution is the supplied hyperlink count. Suspicious-token count, represented host length and dot count are also prominent. Safe-anchor and path measurements contribute, showing that the selected decision uses both URL and webpage signals. These are model-specific associations. A low hyperlink count may characterise some phishing examples and also legitimate sparse pages; it cannot be treated as a sufficient phishing rule.")
local=json.loads((ROOT/'experiments/page_20261007_02/analysis/local_explanations.json').read_text())
table('Table 4-5. Four retained local diagnostic cases at cutoff 0.714698.', ['Published outcome','Score','Top contribution','Log-odds contribution'], [[r['case'].replace('_',' ').title(),val(r['score']),r['top_contributions'][0]['feature'],f"{r['top_contributions'][0]['shap_log_odds']:+.3f}"] for r in local], [1.6,.7,2.0,1.4])
p("In the retained false-positive case, hyperlink count 15 contributes +1.393 log odds and suffix length 2 contributes +1.030; other features partly counteract these values, but the final score remains 0.8032. The legitimate label is therefore missed by the warning decision. In the false-negative case, absence of selected suspicious tokens contributes -0.714 despite other positive evidence, and the score 0.4124 falls below the cutoff. These cases explain actual errors instead of presenting only convincing phishing examples.")
abl=pd.read_csv(ROOT/'experiments/page_20261007_02/analysis/feature_group_ablations.csv')
table('Table 4-6. Exploratory feature-group diagnostics on the same page-study roles, seed 42.', ['Inputs','FPR %','Recall %','F1','AP'], [[r.features.replace('_',' '),pct(r.fpr),pct(r.recall),val(r.f1),val(r.average_precision)] for _,r in abl.iterrows()], [1.65,1.0,1.0,1.0,1.05])
p("The combined model recalls 82.22% compared with 61.08% for URL-only inputs and 56.37% for content-only inputs on these same page-study roles. Their test FPRs are 6.24%, 6.91% and 5.40%, respectively, after independent calibration. This supports a complementary-feature interpretation within the observed partition. The ablations are fitted after primary test inspection and are exploratory; they are not used to substitute a new best architecture or claim superiority on a new blind evaluation.")
h('4.6 Adaptive Replay and Final Checkpoints')
r=pd.read_csv(ROOT/'experiments/page_20261007_02/artifacts/replay_metrics.csv');t=pd.read_csv(ROOT/'experiments/page_20261007_02/artifacts/final_test_metrics.csv')
rg=r.groupby(['order','delay','policy'])[['fpr','recall','f1','updates']].mean();tg=t.groupby(['order','delay','policy'])[['fpr','recall','f1']].mean()
table('Table 4-7. Original replay predictions in path-profile order, mean over three seeds.', ['Delay','Policy','FPR %','Recall %','F1','Updates'], [[d,pol,pct(row.fpr),pct(row.recall),val(row.f1),f'{row.updates:.1f}'] for (order,d,pol),row in rg.iterrows() if order=='path_profile'], [.55,1.25,.9,1.0,1.0,1.0])
p("The fixed page model has replay FPR 3.50% and recall 82.98%. Under immediate labels, threshold-only updating increases mean recall but also raises false alarms, while periodic refitting lowers mean recall. ADWIN makes one update per realization in the path-profile order, giving a small change in original-prediction metrics. With two-batch delay it makes no update. The detector's shorter available error history and decision-version boundary are part of that execution, not errors to be hidden by counting future labels.")
fig('figure_4_4_replay.png','Figure 4-4. Batch-level original prediction error in path-profile order with two-batch delay, seed 42. Coincident fixed and drift traces reflect no drift update.')
p("The replay plot preserves errors at their original arrival decisions. Replacing those predictions with a final model's retrospective scores would overstate online performance. The final partial batch has fewer records and is still included in the aggregate with its actual row count. All 48 realizations, thresholds, label arrivals, update events and final models remain archived.")
table('Table 4-8. Final checkpoint tests in path-profile order on the unchanged 6,849 test rows, mean over three seeds.', ['Delay','Policy','FPR %','Recall %','F1'], [[d,pol,pct(row.fpr),pct(row.recall),val(row.f1)] for (order,d,pol),row in tg.iterrows() if order=='path_profile'], [.6,1.6,1.15,1.2,1.15])
p("With two-batch delay, periodic final checkpoints have mean FPR 4.95% and recall 79.59%, compared with 6.24% and 82.22% for the fixed model. The mean falls below 5% FPR with recall above 70%, but two of the three individual periodic runs have FPR slightly above 5%. Their FPR range is 4.23%-5.37% and recall range 78.25%-80.58%. The mean is not a guarantee for every seed, source or application.")
p("The seed-42 paired whole-domain bootstrap estimates periodic-minus-fixed FPR change between -2.14 and +0.06 percentage points, and recall change between -4.22 and -0.47 percentage points. The false-alarm interval includes no improvement, whereas the recall interval supports a detection loss conditional on these models. Consequently, a general claim of significant adaptive superiority is not supported. Periodic updating is an implemented tradeoff that can help an observed mean constraint at a cost.")
p("Immediate ADWIN final checkpoints have mean FPR 5.94% and recall 82.38%, but one realization reaches FPR 9.54%. Its lower update count therefore does not ensure a reliable final cutoff. Threshold-only final models have still higher false alarms in several conditions. High F1 alone could obscure that behaviour. Policy names, label delay and update counts are reported together so a favourable single checkpoint is not selected as the headline detector.")
table('Table 4-9. Stationary-order final checkpoint tests, mean over three seeds.', ['Delay','Policy','FPR %','Recall %','F1','Updates'], [[d,pol,pct(row.fpr),pct(row.recall),val(row.f1),f'{rg.loc[(order,d,pol),"updates"]:.1f}'] for (order,d,pol),row in tg.iterrows() if order=='stationary'], [.55,1.25,.9,1.0,1.0,1.0])
p("ADWIN performs no updates in either shuffled-delay condition. Periodic and threshold-only policies still update by design and change final predictions even without a constructed profile ordering. This provides a useful unnecessary-update comparator, but the small stream and unavailable natural timestamps limit conclusions about long-run alarm rates. The order control does not demonstrate adversarial drift resistance.")
fig('figure_4_5_tradeoff.png','Figure 4-5. Mean final checkpoint FPR and recall for all four order/delay conditions. Coincident points are retained; the dashed line marks 5% FPR.')
h('4.7 Acceptance and Remaining Risks')
p("The archived-page prototype now provides genuine fitted predictions, reconstructing explanations and executable update policies. It no longer depends on reporting only a failed source model. The fixed detector detects most labelled phishing in its isolated domain test and has useful ranking, while the adaptive comparison shows real constraints rather than assumed improvement. The retained URL transfer failure still prevents a claim that the raw-string model generalises reliably across sources.")
p("The main remaining risks are historical label validity, one-corpus content generalisation, feature-extraction cost and availability, small calibration clusters and finite replay length. The webpage features may carry their creator's collection conventions. The original page measurements have not been independently re-extracted, and present sites are not tested. A second independently measured content archive and a natural-time label stream would be required before claiming deployment readiness.")
endchapter(4)

h('Chapter 5. Conclusion',1)
h('5.1 Summary of the Project')
p("This project implements an explainable phishing prototype and evaluates adaptive decisions using real public archived data. It first establishes that near-perfect within-corpus URL performance can coexist with almost complete external false alarms. The response is a diagnosis, a corrected learning representation, an additional independently published archive and a separate webpage-feature implementation. The final study includes a usable offline record interface, frozen-domain tests, verified explanations and 48 complete policy replays rather than stopping at the original failure.")
p("The revised URL-only branch remains limited. Scheme and root-path shortcuts are corrected, but its transferred false-positive rate and recall do not meet the intended quality. The archived-page branch adds 24 supplied content measurements to 44 recomputed numerical URL inputs. Its selected histogram booster achieves F1 0.8725 and AP 0.9528 on 6,849 domain-isolated test records, detecting 82.22% of published phishing while falsely warning on 6.24% of legitimate records. This is evidence for the archived-input contract within that corpus, not a general web safety guarantee.")
h('5.2 Answers to the Research Questions')
p("RQ1 is answered by separating collection transfer from new-domain detection. The original forest's source score is misleading for cross-source use, and corrected URL features alone do not solve the observed transfer problem. With archived content and corpus-specific fitting, the prototype obtains substantial detection on unseen domains. Its false-positive rate remains above the fixed-model 5% target. A comparison of different corpora and input contracts cannot isolate a single causal improvement; same-partition exploratory ablations provide the narrower evidence for combining URL and content.")
p("RQ2 is answered by numerical verification and error cases. Training-path TreeSHAP reconstructs the selected model's log odds to near floating-point precision on sampled records and identifies both URL and content contributions. A rejected interventional run demonstrates why explanation success cannot be assumed. Local explanations include false positives and false negatives; they expose model evidence without validating the label or creating a causal phishing rule. No evidence of improved human understanding or trust is claimed.")
p("RQ3 is answered by controlled update comparisons. Refitting and recalibration change later decisions, but their benefit depends on order, feedback delay and cutoff variation. Periodic updating under two-batch delay in path-profile order lowers mean final FPR to 4.95% and lowers recall to 79.59%. The paired uncertainty check does not establish a general false-alarm improvement and does support a recall cost for the seed-42 contrast. ADWIN avoids updates in the shuffled control but sometimes produces a high-FPR final checkpoint after an alarm. Adaptivity is therefore implemented and measured, with conditions and tradeoffs.")
h('5.3 Contributions and Practical Use')
p("The practical contribution is an inspectable offline software pipeline. A reviewer can follow each dataset from its checksum to a canonical row, domain role, fitted feature matrix, checkpoint and original prediction. The public record interface validates the required webpage inputs and returns a model decision together with its cutoff, decision version and signed explanations. The label ledger prevents use of future labels and prevents older decision errors from entering a reset detector. The numerical verification checks the stored evidence independently.")
p("The methodological contribution is the connection between diagnosis and revised implementation. Source-specific cues are investigated by actual perturbations, thresholds are evaluated with both legitimate and phishing denominators, and page-domain performance is distinguished from cross-source transfer. All policies and failed configurations remain available. These are useful engineering and evaluation outcomes for the project, while the classifier, SHAP method and ADWIN algorithm themselves are established methods.")
h('5.4 Limitations')
p("The datasets contain historical observations with adopted creator labels. No record is relabelled through a present-day site visit. Excluding malformed URLs and removing duplicate/conflicting keys changes the studied population; those exclusions are quantified but cannot make every original label reliable. Registered-domain isolation reduces one form of leakage without excluding common templates, campaigns or public-feed overlap.")
p("The new page result uses one fixed grouped partition from one content archive. Initial boosting fits are deterministic across the listed seeds, and replay seed variation does not supply independent datasets. Whole-domain intervals capture finite-domain sampling conditional on the fitted models; they omit measurement error, alternate data collection and model-selection uncertainty. The 555-row calibration sample cannot guarantee a population FPR budget, as the observed 6.24% test rate illustrates.")
p("The webpage measurements are supplied in the archive and have not been recomputed from original HTML. The program consequently accepts structured archived records and cannot presently derive the required measurements from arbitrary online URLs. Some feature groups have correlated meanings and creator-specific scales. Excluding external-service inputs avoids a direct dependence on reputation and indexing services, but does not remove every collection convention from the remaining features.")
p("The replay is short and its order is constructed. Path length profiles are not authentic observation times, and two-batch latency is not a calibrated real-world verification delay. The ADWIN experiment monitors errors of the current decision version, which limits available feedback after updates. It does not measure natural drift frequency or adversarial attacks. The empirical policy comparison is complete for its controlled protocol but insufficient for long-term operation.")
h('5.5 Future Work')
p("The next research requirement is an untouched second corpus containing independently documented webpage measurements with the same input meanings. Its domain and campaign overlap should be checked before fitting or tuning. Prospective collection should retain actual observation and label-verification times so delayed adaptation can be studied chronologically. New redesign using this study's inspected holdout must be labelled exploratory until new evidence is collected.")
p("For the software, a sandboxed content extractor would need independent validation against the archive's feature definitions, with explicit handling of missing pages and failed requests. Cutoff updates should be studied with larger calibration-domain sets and uncertainty-aware guardrails. These are future experiments, not implemented results. They should be accepted based on low legitimate warning rates and retained phishing recall together.")
p("Explanation evaluation should include user tasks and failure comprehension rather than numerical reconstruction alone. More partitions and longer controlled streams could assess stability, update cost and unnecessary alarms. The present artifacts provide a reproducible starting point for those investigations while already delivering a functioning archived-page detector and a complete, evidence-backed adaptive comparison.")
endchapter(5)
# The original Chapter-3 Cite controls must each be transferred exactly once.
assert all(not cites[i] for i in [129,132,137,147,149,153,158,160])
# Preserve ThesisBody bookmark start from Chapter 1 and relocate its end before References.
starts=xp(body,'.//w:bookmarkStart[@w:name="ThesisBody"]')
if starts:
    ident=starts[0].get(qn('w:id'))
    for x in list(xp(body,f'.//w:bookmarkEnd[@w:id="{ident}"]')):x.getparent().remove(x)
    last=next(x for x in reversed(new_nodes) if x.tag==qn('w:p') and not xp(x,'.//w:sectPr'))
    e=OxmlElement('w:bookmarkEnd');e.set(qn('w:id'),ident);last.append(e)
# Keep image identifiers unique after adding new figures to a Word-produced package.
for i,x in enumerate(xp(body,'.//wp:docPr'),100):x.set('id',str(i))
new_ids=[tag(x) for x in xp(body,'.//w:sdt') if 'MENDELEY_CITATION' in tag(x)]
assert sorted(old_ids)==sorted(new_ids) and len(new_ids)==30
# New source controls are inserted later by the actual Mendeley Cite add-in.
doc.core_properties.author='Geoff';doc.core_properties.title='Explainable and Adaptive Machine Learning for Phishing Detection'
doc.save(FINAL)
# Preserve all unrelated package parts byte-for-byte from the actual linked document.
allowed={'word/document.xml','word/_rels/document.xml.rels','docProps/core.xml','docProps/app.xml','[Content_Types].xml'}
with ZipFile(BASE) as z:prior={n:z.read(n) for n in z.namelist()}
with ZipFile(FINAL) as z:out={n:z.read(n) for n in z.namelist()}
for n,blob in prior.items():
    if n not in allowed and n in out:out[n]=blob
with ZipFile(FINAL,'w',ZIP_DEFLATED) as z:
    for n,blob in out.items():z.writestr(n,blob)
(BUILD/'chapter_blocks.json').write_text(json.dumps(blocks,indent=2)+'\n')
(BUILD/'rewrite_audit.json').write_text(json.dumps({'retained_real_cite_controls':30,'new_sources_pending_actual_addin':3,'prior_sha256':hashlib.sha256(BASE.read_bytes()).hexdigest(),'rewritten_chapters':[3,4,5],'new_headings':heads,'unchanged_parts_restored':True},indent=2)+'\n')
md=['# Explainable and Adaptive Machine Learning for Phishing Detection','','Author Geoff. Revised 7 October 2026.','']
for x in body:
    if x.tag==qn('w:p'):
        t=text(x)
        if t:md.extend([t,''])
(ROOT/'thesis/research_manuscript.md').write_text('\n'.join(md))
print('Rewritten DOCX with genuine retained controls:',len(new_ids),'new chapter blocks',len(blocks))
