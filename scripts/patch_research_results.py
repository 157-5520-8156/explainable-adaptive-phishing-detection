"""Update corrected comparator evidence while preserving real Cite controls."""
from pathlib import Path
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree as E
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
import pandas as pd,json,re
ROOT=Path(__file__).resolve().parents[1];path=ROOT/'thesis/Phishing_Detection_Mendeley_Draft.docx'
with ZipFile(path) as z:prior={n:z.read(n) for n in z.namelist()}
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
r=E.fromstring(prior['word/document.xml']);tags=r.xpath('//w:sdtPr/w:tag/@w:val',namespaces=ns)
d=Document(path);A=ROOT/'experiments/page_20261007_02/artifacts'
replay=pd.read_csv(A/'replay_metrics.csv').groupby(['order','delay','policy'])[['fpr','recall','f1','updates']].mean()
final=pd.read_csv(A/'final_test_metrics.csv').groupby(['order','delay','policy'])[['fpr','recall','f1']].mean()
def pct(v):return f'{100*float(v):.2f}'
def val(v):return f'{float(v):.4f}'
rows={
'Table 4-7.':[[d,p,pct(v.fpr),pct(v.recall),val(v.f1),f'{v.updates:.1f}'] for (o,d,p),v in replay.iterrows() if o=='path_profile'],
'Table 4-8.':[[d,p,pct(v.fpr),pct(v.recall),val(v.f1)] for (o,d,p),v in final.iterrows() if o=='path_profile'],
'Table 4-9.':[[d,p,pct(v.fpr),pct(v.recall),val(v.f1),f'{replay.loc[(o,d,p),"updates"]:.1f}'] for (o,d,p),v in final.iterrows() if o=='stationary']}
from docx.table import Table
for p in d.paragraphs:
 for prefix,values in rows.items():
  if p.text.startswith(prefix):
   el=p._p.getnext();assert el.tag==qn('w:tbl');t=Table(el,d._body)
   for row,items in zip(t.rows[1:],values):
    for c,item in zip(row.cells,items):
     c.paragraphs[0].runs[0].text=str(item)
# Rewrite only plain paragraphs; no Cite SDT inside these slots.
def replace(p,s):
 assert not p._p.xpath('.//w:sdt');p.text=s
 for run in p.runs:run.font.name='Arial';run.font.size=Pt(11)
for p in d.paragraphs:
 if p.text.startswith('The update pool combines'):
  # Contains retained group-split Cite; leave the control untouched and change plain runs only.
  for run in p.runs:
   run.text=run.text.replace('The threshold-only policy leaves the estimator fixed.', 'The threshold-only policy leaves the estimator fixed and uses only original calibration plus revealed recent replay rows, excluding every original fitting domain.')
 if p.text.startswith('The project now contains'):pass
 if p.text.startswith('A saved real record'):
  replace(p,p.text.replace('Eight behaviour tests pass.','Nine behaviour tests pass.'))
 if p.text.startswith('The mean satisfies the study'):
  pass
 if p.text.startswith('With two-batch delay, periodic final checkpoints'):
  replace(p,p.text.replace("The mean satisfies the study's final adaptive FPR/recall targets of 5% and 70%,",'The mean falls below 5% FPR with recall above 70%,').replace('but two of','but two of'))
 if p.text.startswith('The fixed page model has replay'):
  replace(p,p.text.replace('threshold-only updating increases mean recall but also raises false alarms','corrected threshold-only updating raises mean recall to 85.12% and FPR to 4.55%'))
 if p.text.startswith('Immediate ADWIN final checkpoints'):
  extra=' The first threshold comparator reused some domains from its unchanged estimator fitting pool during calibration and was withdrawn from comparison. The corrected calibration uses original calibration and revealed replay domains only. Its 12 reruns occur after test inspection and are correction diagnostics, not new blind evidence. The 36 unaffected executions and their frozen tests are reused unchanged; the original contaminated archive remains inspectable.'
  replace(p,p.text+extra)
# Replace image pixels while retaining package relationships, geometry and captions.
from docx.oxml.ns import nsmap
for p in d.paragraphs:
 if p.text.startswith('Figure 4-4.') or p.text.startswith('Figure 4-5.'):
  image_node=p._p.getprevious();ids=image_node.xpath('.//a:blip/@r:embed')
  assert len(ids)==1
  part=d.part.related_parts[ids[0]]
  filename='figure_4_4_replay.png' if p.text.startswith('Figure 4-4.') else 'figure_4_5_tradeoff.png'
  part._blob=(ROOT/'experiments/page_20261007_02/analysis/figures'/filename).read_bytes()
d.save(path)
with ZipFile(path) as z:out={n:z.read(n) for n in z.namelist()}
# Restore every unrelated part exactly; only document text, image bytes and app/core metadata change.
for n,blob in prior.items():
 if n not in {'word/document.xml','docProps/core.xml','docProps/app.xml'} and not n.startswith('word/media/'):
  out[n]=blob
new=E.fromstring(out['word/document.xml']);assert tags==new.xpath('//w:sdtPr/w:tag/@w:val',namespaces=ns)
with ZipFile(path,'w',ZIP_DEFLATED) as z:
 for n,blob in out.items():z.writestr(n,blob)
# The authoring recipe also points to the accepted correction archive for future reproducibility.
recipe=ROOT/'scripts/write_research_revision.py';s=recipe.read_text().replace('page_20261007_01','page_20261007_02').replace('Eight behaviour tests pass.','Nine behaviour tests pass.')
recipe.write_text(s)
print('Updated corrected comparator tables and figures; all authentic Cite tags preserved')
